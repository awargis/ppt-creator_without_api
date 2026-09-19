import io
from pptx import Presentation
from .template_validator import validate_template
from .layout_engine import fit_box

def build_slides(template_bytes, regions, answers):
    prs = validate_template(template_bytes)
    layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
    while len(prs.slides): prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])
    for region in regions:
        slide = prs.slides.add_slide(layout)
        stream = io.BytesIO(); region.image.save(stream, "PNG"); stream.seek(0)
        left, top, width, height = fit_box(prs, region.image)
        slide.shapes.add_picture(stream, left, top, width=width, height=height)
        box = slide.shapes.add_textbox(0, 0, prs.slide_width, 500000)
        box.text_frame.text = f"Q{region.number}  •  {region.subject}    Answer: {answers.get(region.number, '—')}"
    return prs
