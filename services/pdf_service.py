import copy
import io
from pptx import Presentation
from pptx.util import Emu

RELATIONSHIP_NAMESPACE = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

def duplicate_slide(prs, source_index):
    source = prs.slides[source_index]
    new_slide = prs.slides.add_slide(source.slide_layout)
    
    # Strip layout placeholders to prevent ghost shapes/stacking
    for shape in list(new_slide.shapes): 
        shape._element.getparent().remove(shape._element)
    
    rel_map = {}
    for old_id, rel in source.part.rels.items():
        if "notesSlide" in rel.reltype or "slideLayout" in rel.reltype: continue
        rel_map[old_id] = new_slide.part.relate_to(
            rel.target_ref if rel.is_external else rel.target_part, 
            rel.reltype, 
            is_external=rel.is_external
        )
        
    for shape in source.shapes:
        copied_el = copy.deepcopy(shape._element)
        for el in copied_el.iter():
            for attr, val in list(el.attrib.items()):
                if attr.startswith(RELATIONSHIP_NAMESPACE) and val in rel_map: 
                    el.set(attr, rel_map[val])
        new_slide.shapes._spTree.append(copied_el)
    return new_slide

def build_subject_ppt(template_bytes: bytes, questions: list[dict], answers: dict) -> bytes:
    prs = Presentation(io.BytesIO(template_bytes))
    
    # Standard Vidyapeeth content coordinates
    c_left, c_top = Emu(int(0.45 * 914400)), Emu(int(1.50 * 914400))
    c_width = prs.slide_width - Emu(int(0.9 * 914400))
    c_height = prs.slide_height - c_top - Emu(int(0.35 * 914400))

    for q in questions:
        # STRICT: 1 Question = 1 Fresh Slide duplicated from Master (Index 0)
        new_slide = duplicate_slide(prs, 0)
        
        # Hard-replace text elements for Question Number & Answer Key
        for shape in new_slide.shapes:
            if shape.has_text_frame:
                text = shape.text
                if "#QUESTION" in text:
                    shape.text_frame.text = text.replace("#QUESTION", f"Q{q['number']}")
                elif "Ans." in text:
                    ans_val = answers.get(q['number'], answers.get(str(q['number']), " "))
                    shape.text_frame.text = f"Ans. ({ans_val})"

        # Scale and inject dark mode image without distorting aspect ratio
        img_buf = io.BytesIO()
        q["image"].save(img_buf, format="PNG")
        img_buf.seek(0)
        
        iw, ih = q["image"].size
        scale = min(int(c_width) / iw, int(c_height) / ih)
        fw, fh = max(1, int(iw * scale)), max(1, int(ih * scale))
        
        # Center the image precisely
        left = int(c_left) + (int(c_width) - fw) // 2
        top = int(c_top) + (int(c_height) - fh) // 2
        
        new_slide.shapes.add_picture(img_buf, left, top, width=fw, height=fh)

    # Delete the master slide so the final presentation starts immediately with Q1
    xml_slides = prs.slides._sldIdLst
    xml_slides.remove(xml_slides[0])
    
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
