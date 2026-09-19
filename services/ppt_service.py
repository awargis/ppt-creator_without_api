import io
import copy
import re
from pptx import Presentation
from pptx.util import Inches

REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def duplicate_slide(prs, source_index=0):
    source = prs.slides[source_index]
    target = prs.slides.add_slide(source.slide_layout)
    for shape in list(target.shapes):
        target.shapes._spTree.remove(shape._element)
    rels = {}
    for rid, rel in source.part.rels.items():
        if "notesSlide" in rel.reltype or "slideLayout" in rel.reltype:
            continue
        rels[rid] = target.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,
                                          rel.reltype, is_external=rel.is_external)
    for shape in source.shapes:
        element = copy.deepcopy(shape._element)
        for child in element.iter():
            for key, value in list(child.attrib.items()):
                if key.startswith(REL_NS) and value in rels:
                    child.set(key, rels[value])
        target.shapes._spTree.insert_element_before(element, "p:extLst")
    return target


def _replace_tokens(slide, number, answer):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace("#QUESTION", f"Q{number}")
                run.text = run.text.replace("#ANSWER", answer or "—")
                run.text = run.text.replace("#SUBJECT", "")


def build_subject_ppt(template_bytes: bytes, questions: list[dict], answers: dict, style: str = "Premium Light") -> bytes:
    prs = Presentation(io.BytesIO(template_bytes))
    if not prs.slides:
        raise ValueError("The template must contain at least one slide")
    master_index = 0
    content_left, content_top = Inches(0.45), Inches(1.25)
    content_width = prs.slide_width - Inches(0.9)
    content_height = prs.slide_height - Inches(1.55)
    for question in questions:
        slide = duplicate_slide(prs, master_index)
        answer = str(answers.get(question["number"], ""))
        _replace_tokens(slide, question["number"], answer)
        image = question["image"]
        stream = io.BytesIO()
        image.save(stream, format="PNG", optimize=True)
        stream.seek(0)
        iw, ih = image.size
        scale = min(int(content_width) / iw, int(content_height) / ih)
        width, height = max(1, int(iw * scale)), max(1, int(ih * scale))
        left = int(content_left) + (int(content_width) - width) // 2
        top = int(content_top) + (int(content_height) - height) // 2
        slide.shapes.add_picture(stream, left, top, width=width, height=height)
    ids = prs.slides._sldIdLst
    ids.remove(ids[0])
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()
